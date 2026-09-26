"""Study-focused source validation for the TARGET-AML survival analysis.

The goal is not to inventory the entire GDC schema. It checks whether the
public sources can support the planned analysis and records the ambiguities
that must be resolved before cohort construction.
"""

from __future__ import annotations

import argparse
import csv
import json
from collections import Counter
from collections.abc import Mapping, Sequence
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import requests

from pediastat.config import PROJECT_ROOT
from pediastat.gdc_api import (
    GDC_API_BASE_URL,
    USER_AGENT,
    classify_vital_status,
    download_open_file,
    fetch_all_cases,
    fetch_clinical_files,
    fetch_project,
    flatten_clinical_file_row,
    md5_file,
)

DEFAULT_PAGE_SIZE = 500
AUDIT_TIMEOUT_SECONDS = 120.0
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "artifacts" / "source_audit"
DEFAULT_DOWNLOAD_DIR = PROJECT_ROOT / "data" / "raw" / "gdc_open_clinical_supplements"

# Only fields that motivate the population, endpoint, or baseline analysis.
AUDIT_FIELDS = (
    "case_id",
    "submitter_id",
    "demographic.vital_status",
    "demographic.days_to_death",
    "demographic.sex_at_birth",
    "diagnoses.age_at_diagnosis",
    "diagnoses.days_to_last_follow_up",
    "follow_ups.days_to_follow_up",
)

MISSING_TOKENS = frozenset(
    "not reported|unknown|not allowed to collect|not applicable|missing|"
    "not available|--|n/a|na".split("|")
)

FIELD_COLUMNS = (
    "field_path",
    "n_cases",
    "n_available",
    "pct_missing",
    "n_usable",
    "n_cases_with_multiple_values",
    "n_cases_with_disagreeing_values",
    "numeric_min",
    "numeric_max",
)

FILE_COLUMNS = (
    "file_id",
    "file_name",
    "data_category",
    "data_type",
    "data_format",
    "access",
    "file_size",
    "md5sum",
    "state",
    "associated_project",
    "n_associated_cases",
)


def _now() -> str:
    return datetime.now(UTC).replace(microsecond=0).isoformat()


def _is_null(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _is_usable(value: Any) -> bool:
    return not _is_null(value) and not (
        isinstance(value, str) and value.strip().lower() in MISSING_TOKENS
    )


def _number(value: Any) -> float | None:
    if isinstance(value, bool) or not isinstance(value, int | float):
        return None
    number = float(value)
    return number if number == number else None


def get_values_at_path(record: Mapping[str, Any], path: str) -> list[Any]:
    """Return every value at a dotted path; never pick an arbitrary nested row."""
    values: list[Any] = [record]
    for part in path.split("."):
        next_values = []
        for value in values:
            for item in value if isinstance(value, list) else [value]:
                if isinstance(item, Mapping) and part in item:
                    next_values.append(item[part])
        values = next_values
    return values


def summarize_field(
    cases: Sequence[Mapping[str, Any]],
    field_path: str,
) -> dict[str, Any]:
    """Summarize availability and repeated-value disagreement for one study field."""
    available = usable = multiple = disagree = 0
    numbers: list[float] = []

    for case in cases:
        values = [v for v in get_values_at_path(case, field_path) if not _is_null(v)]
        available += bool(values)
        usable += any(_is_usable(v) for v in values)
        if len(values) > 1:
            multiple += 1
            disagree += len({repr(v) for v in values}) > 1
        numbers.extend(n for v in values if (n := _number(v)) is not None)

    n_cases = len(cases)
    missing = n_cases - available
    return {
        "field_path": field_path,
        "n_cases": n_cases,
        "n_available": available,
        "pct_missing": round(missing / n_cases * 100, 2) if n_cases else 0.0,
        "n_usable": usable,
        "n_cases_with_multiple_values": multiple,
        "n_cases_with_disagreeing_values": disagree,
        "numeric_min": min(numbers) if numbers else None,
        "numeric_max": max(numbers) if numbers else None,
    }


def _numeric_values(case: Mapping[str, Any], path: str) -> list[float]:
    return [
        number
        for value in get_values_at_path(case, path)
        if (number := _number(value)) is not None
    ]


def audit_survival_fields(cases: Sequence[Mapping[str, Any]]) -> dict[str, Any]:
    """Compare plausible OS constructions before the endpoint is locked."""
    status_counts: Counter[str] = Counter()
    c: Counter[str] = Counter()

    for case in cases:
        raw_status = next(
            (
                value
                for value in get_values_at_path(case, "demographic.vital_status")
                if not _is_null(value)
            ),
            None,
        )
        status = classify_vital_status(raw_status)
        status_counts[status] += 1

        death = _numeric_values(case, "demographic.days_to_death")
        dx_fu = _numeric_values(case, "diagnoses.days_to_last_follow_up")
        fu = _numeric_values(case, "follow_ups.days_to_follow_up")

        c["negative_death"] += sum(x < 0 for x in death)
        c["negative_dx_fu"] += sum(x < 0 for x in dx_fu)
        c["negative_fu"] += sum(x < 0 for x in fu)
        c["zero_death"] += sum(x == 0 for x in death)
        c["zero_dx_fu"] += sum(x == 0 for x in dx_fu)
        c["zero_fu"] += sum(x == 0 for x in fu)
        c["multiple_fu"] += len(fu) > 1
        c["dx_vs_fu_disagree"] += bool(dx_fu and fu and max(dx_fu) != max(fu))

        if status == "dead":
            if any(x >= 0 for x in death):
                c["dead_valid"] += 1
            else:
                c["dead_missing_time"] += 1
        elif status == "alive":
            if any(x >= 0 for x in dx_fu):
                c["alive_dx_valid"] += 1
            if any(x >= 0 for x in fu):
                c["alive_fu_valid"] += 1
            if any(x >= 0 for x in dx_fu) or any(x >= 0 for x in fu):
                c["alive_either_valid"] += 1
            else:
                c["alive_missing_time"] += 1

    gdc_style = c["dead_valid"] + c["alive_either_valid"]
    diagnosis_only = c["dead_valid"] + c["alive_dx_valid"]
    follow_up_only = c["dead_valid"] + c["alive_fu_valid"]

    return {
        "n_cases": len(cases),
        "vital_status_class_counts": dict(status_counts),
        "n_dead_missing_days_to_death": c["dead_missing_time"],
        "n_alive_missing_all_follow_up_times": c["alive_missing_time"],
        "n_cases_with_multiple_follow_up_times": c["multiple_fu"],
        "n_cases_diagnosis_vs_follow_up_time_disagree": c["dx_vs_fu_disagree"],
        "n_negative_days_to_death_values": c["negative_death"],
        "n_negative_days_to_last_follow_up_values": c["negative_dx_fu"],
        "n_negative_days_to_follow_up_values": c["negative_fu"],
        "n_zero_days_to_death_values": c["zero_death"],
        "n_zero_days_to_last_follow_up_values": c["zero_dx_fu"],
        "n_zero_days_to_follow_up_values": c["zero_fu"],
        "n_cases_with_possible_gdc_style_survival_time": gdc_style,
        "candidate_os_counts": {
            "death_plus_either_follow_up": gdc_style,
            "death_plus_diagnosis_follow_up": diagnosis_only,
            "death_plus_follow_up_records": follow_up_only,
        },
        "unresolved": [
            "Which follow-up source should define censoring time?",
            "How should disagreement between follow-up sources be handled?",
            "Are zero-day times valid or data errors?",
            "Unknown/missing vital status must not be treated as censored.",
        ],
    }


def _xlsx_sheets(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        output = []
        for name in workbook.sheetnames:
            rows = workbook[name].iter_rows(values_only=True)
            header = next(rows, None) or ()
            output.append(
                {
                    "sheet": name,
                    "n_rows": sum(1 for _ in rows),
                    "columns": [
                        "" if value is None else str(value) for value in header
                    ],
                }
            )
        return output
    finally:
        workbook.close()


def _download_supplements(
    rows: Sequence[Mapping[str, Any]],
    download_dir: Path,
    *,
    session: requests.Session,
    timeout: float,
) -> list[dict[str, Any]]:
    """Download only open clinical supplements and inspect workbook columns."""
    inspections = []

    for row in rows:
        file_id = str(row.get("file_id") or "")
        file_name = str(row.get("file_name") or file_id)
        record: dict[str, Any] = {
            "file_id": file_id,
            "file_name": file_name,
            "access": row.get("access"),
        }

        if str(row.get("access") or "").lower() != "open":
            record["downloaded"] = False
            inspections.append(record)
            continue

        destination = download_dir / file_name
        download_open_file(
            file_id,
            destination,
            session=session,
            timeout=timeout,
        )
        checksum = md5_file(destination)
        record.update(
            downloaded=True,
            checksum_match=(
                checksum == row.get("md5sum") if row.get("md5sum") else None
            ),
        )
        if destination.suffix.lower() == ".xlsx":
            record["sheets"] = _xlsx_sheets(destination)
        inspections.append(record)

    return inspections


def _write_json(path: Path, value: Any) -> None:
    path.write_text(
        json.dumps(value, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def _write_csv(
    path: Path,
    columns: Sequence[str],
    rows: Sequence[Mapping[str, Any]],
) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle, fieldnames=list(columns), extrasaction="ignore"
        )
        writer.writeheader()
        writer.writerows(rows)


def run_audit(
    *,
    output_dir: Path = DEFAULT_OUTPUT_DIR,
    download_dir: Path = DEFAULT_DOWNLOAD_DIR,
    skip_download: bool = False,
    page_size: int = DEFAULT_PAGE_SIZE,
    timeout: float = AUDIT_TIMEOUT_SECONDS,
) -> dict[str, Any]:
    """Validate that the public sources can support the planned survival study."""
    session = requests.Session()
    session.headers["User-Agent"] = USER_AGENT

    try:
        started = _now()
        project = fetch_project(timeout=timeout, session=session)
        cases, pagination = fetch_all_cases(
            AUDIT_FIELDS,
            page_size=page_size,
            timeout=timeout,
            session=session,
        )
        fields = [summarize_field(cases, field) for field in AUDIT_FIELDS]
        survival = audit_survival_fields(cases)
        files = [
            flatten_clinical_file_row(row)
            for row in fetch_clinical_files(timeout=timeout, session=session)
        ]
        supplements = (
            []
            if skip_download
            else _download_supplements(
                files,
                download_dir,
                session=session,
                timeout=timeout,
            )
        )

        metadata = {
            "audit_timestamp_utc": started,
            "api_base_url": GDC_API_BASE_URL,
            "project_id": project.get("project_id"),
            "name": project.get("name"),
            "n_cases_retrieved": len(cases),
            "cases_api_pagination_total": pagination.get("total"),
            "question": (
                "Can public TARGET-AML clinical sources support a defensible "
                "pediatric overall-survival analysis?"
            ),
        }

        output_dir.mkdir(parents=True, exist_ok=True)
        _write_json(output_dir / "project_metadata.json", metadata)
        _write_csv(
            output_dir / "clinical_field_availability.csv",
            FIELD_COLUMNS,
            fields,
        )
        _write_json(output_dir / "survival_field_audit.json", survival)
        _write_csv(output_dir / "open_clinical_files.csv", FILE_COLUMNS, files)
        if supplements:
            _write_json(
                output_dir / "open_clinical_supplement_columns.json",
                supplements,
            )

        # This generic schema artifact belonged to the old audit design.
        (output_dir / "entity_cardinality.json").unlink(missing_ok=True)

        return {
            "project_metadata": metadata,
            "field_summaries": fields,
            "survival": survival,
            "clinical_files": files,
            "supplement_inspections": supplements,
            "n_cases": len(cases),
            "output_dir": str(output_dir),
        }
    finally:
        session.close()


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Validate public TARGET-AML sources for the survival analysis."
    )
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--download-dir", type=Path, default=DEFAULT_DOWNLOAD_DIR)
    parser.add_argument("--skip-download", action="store_true")
    parser.add_argument("--page-size", type=int, default=DEFAULT_PAGE_SIZE)
    args = parser.parse_args(argv)

    result = run_audit(
        output_dir=args.output_dir,
        download_dir=args.download_dir,
        skip_download=args.skip_download,
        page_size=args.page_size,
    )
    print("TARGET-AML source validation")
    print(f"Cases retrieved: {result['n_cases']}")
    print(
        "Possible GDC-style OS times: "
        f"{result['survival']['n_cases_with_possible_gdc_style_survival_time']}"
    )
    print(f"Artifacts: {result['output_dir']}")
    print("No patient-level dataset was written.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
