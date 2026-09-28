"""Parse public TARGET-AML clinical sources and run source-reconciliation QA.

This module handles source-faithful ingestion and pre-cohort QA:
- TARGET identifier normalization
- missing-value classification
- GDC case/entity parsing
- TARGET clinical supplement workbook parsing and profiling
- source overlap and discordance checks

Database persistence belongs in ``pediastat.database``.
Cohort construction belongs in ``pediastat.cohort``.
"""

from __future__ import annotations

import csv
import json
import re
from collections import Counter, defaultdict
from collections.abc import Mapping, Sequence
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.engine import Engine

from pediastat.config import PROJECT_ROOT
from pediastat.gdc import as_records

CANONICAL_BARCODE = re.compile(r"^TARGET-\d{2}-[A-Z0-9]+$")
EXTENDED_BARCODE = re.compile(r"^TARGET-\d{2}-[A-Z0-9]+(?:-[A-Z0-9]+)+$")
JOIN_BARCODE = re.compile(r"^(TARGET-\d{2}-[A-Z0-9]+)")


def original_identifier(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, str) and value.strip() == "":
        return None
    return str(value)


def normalize_identifier(value: Any) -> str | None:
    original = original_identifier(value)
    return None if original is None else original.strip().upper()


def join_barcode(value: Any) -> str | None:
    normalized = normalize_identifier(value)
    if normalized is None:
        return None
    match = JOIN_BARCODE.match(normalized)
    return match.group(1) if match else normalized


def identifier_shape(value: Any) -> str:
    original = original_identifier(value)
    if original is None:
        return "missing"
    stripped = original.strip()
    if original != stripped:
        return "whitespace"
    normalized = stripped.upper()
    if CANONICAL_BARCODE.match(normalized):
        return "canonical"
    if EXTENDED_BARCODE.match(normalized):
        return "extended"
    return "malformed"


def summarize_identifiers(values: list[Any]) -> dict[str, Any]:
    originals = [original_identifier(value) for value in values]
    present = [value for value in originals if value is not None]
    shapes: Counter[str] = Counter()
    normalized_counts: Counter[str] = Counter()
    whitespace = 0
    case_differences = 0
    for original in present:
        shapes[identifier_shape(original)] += 1
        if original != original.strip():
            whitespace += 1
        stripped = original.strip()
        if stripped != stripped.upper():
            case_differences += 1
        normalized = normalize_identifier(original)
        if normalized is not None:
            normalized_counts[normalized] += 1
    duplicate_counts = [count for count in normalized_counts.values() if count > 1]
    return {
        "n_records": len(values),
        "n_non_null": len(present),
        "n_null": len(values) - len(present),
        "n_unique_normalized": len(normalized_counts),
        "n_duplicated_normalized_ids": len(duplicate_counts),
        "n_duplicate_records": sum(count - 1 for count in duplicate_counts),
        "n_whitespace": whitespace,
        "n_case_differences": case_differences,
        "n_canonical": shapes["canonical"],
        "n_extended": shapes["extended"],
        "n_malformed": shapes["malformed"],
        "shapes": dict(shapes),
    }


NOT_REPORTED = frozenset({"not reported", "notreported"})
UNKNOWN = frozenset(
    {
        "unknown",
        "unspecified",
        "na",
        "n/a",
        "n.a.",
        "n.a",
        "--",
        ".",
        "missing",
        "not available",
        "null",
    }
)
NOT_APPLICABLE = frozenset({"not applicable", "notapplicable"})
SENTINEL_NUMBERS = frozenset({-99, -999, -9999})


def classify_missing(value: Any) -> str:
    if value is None:
        return "structurally_missing"
    if isinstance(value, str):
        stripped = value.strip()
        if not stripped:
            return "structurally_missing"
        token = " ".join(stripped.lower().split())
        compact = token.replace(" ", "")
        if token in NOT_REPORTED or compact in NOT_REPORTED:
            return "not_reported"
        if token in NOT_APPLICABLE or compact in NOT_APPLICABLE:
            return "not_applicable"
        if token in UNKNOWN or compact in UNKNOWN:
            return "unknown"
        return "observed"
    if isinstance(value, bool):
        return "observed"
    if isinstance(value, int | float) and value in SENTINEL_NUMBERS:
        return "sentinel"
    return "observed"


def is_observed(value: Any) -> bool:
    return classify_missing(value) == "observed"


CASE_NESTED_FIELDS = {
    "demographic",
    "diagnoses",
    "follow_ups",
    "samples",
    "aliquot_ids",
    "submitter_aliquot_ids",
    "sample_ids",
    "submitter_sample_ids",
    "diagnosis_ids",
    "submitter_diagnosis_ids",
}


def _case_keys(case_id: Any, submitter_id: Any) -> dict[str, Any]:
    return {
        "case_id": case_id,
        "submitter_id": submitter_id,
        "submitter_id_normalized": normalize_identifier(submitter_id),
        "join_barcode": join_barcode(submitter_id),
    }


def _as_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value).lower()
    return str(value)


def parse_case_entities(case: Mapping[str, Any]) -> dict[str, list[dict[str, Any]]]:
    case_id = case.get("case_id") or case.get("id")
    submitter_id = case.get("submitter_id")
    keys = _case_keys(case_id, submitter_id)

    project = case.get("project")
    project_id = project.get("project_id") if isinstance(project, dict) else None

    case_row = {
        **keys,
        "project_id": project_id,
        "disease_type": case.get("disease_type"),
        "primary_site": case.get("primary_site"),
        "index_date": _as_text(case.get("index_date")),
        "lost_to_followup": _as_text(case.get("lost_to_followup")),
        "days_to_lost_to_followup": case.get("days_to_lost_to_followup"),
        "payload": {
            key: value for key, value in case.items() if key not in CASE_NESTED_FIELDS
        },
    }

    demographics = [
        {
            **keys,
            "demographic_id": item.get("demographic_id"),
            "vital_status": item.get("vital_status"),
            "days_to_death": item.get("days_to_death"),
            "age_at_index": item.get("age_at_index"),
            "days_to_birth": item.get("days_to_birth"),
            "sex_at_birth": item.get("sex_at_birth"),
            "race": item.get("race"),
            "ethnicity": item.get("ethnicity"),
            "year_of_birth": item.get("year_of_birth"),
            "year_of_death": item.get("year_of_death"),
            "cause_of_death": item.get("cause_of_death"),
            "age_is_obfuscated": _as_text(item.get("age_is_obfuscated")),
            "payload": dict(item),
        }
        for item in as_records(case.get("demographic"))
    ]

    diagnoses: list[dict[str, Any]] = []
    treatments: list[dict[str, Any]] = []
    for diagnosis in as_records(case.get("diagnoses")):
        diagnosis_id = diagnosis.get("diagnosis_id")
        diagnoses.append(
            {
                **keys,
                "diagnosis_id": diagnosis_id,
                "age_at_diagnosis": diagnosis.get("age_at_diagnosis"),
                "days_to_diagnosis": diagnosis.get("days_to_diagnosis"),
                "days_to_last_follow_up": diagnosis.get("days_to_last_follow_up"),
                "primary_diagnosis": diagnosis.get("primary_diagnosis"),
                "morphology": diagnosis.get("morphology"),
                "tissue_or_organ_of_origin": diagnosis.get("tissue_or_organ_of_origin"),
                "site_of_resection_or_biopsy": diagnosis.get(
                    "site_of_resection_or_biopsy"
                ),
                "year_of_diagnosis": _as_text(diagnosis.get("year_of_diagnosis")),
                "icd_10_code": diagnosis.get("icd_10_code"),
                "classification_of_tumor": _as_text(
                    diagnosis.get("classification_of_tumor")
                ),
                "diagnosis_is_primary_disease": _as_text(
                    diagnosis.get("diagnosis_is_primary_disease")
                ),
                "payload": {
                    key: value
                    for key, value in diagnosis.items()
                    if key != "treatments"
                },
            }
        )
        for treatment in as_records(diagnosis.get("treatments")):
            treatments.append(
                {
                    **keys,
                    "diagnosis_id": diagnosis_id,
                    "treatment_id": treatment.get("treatment_id"),
                    "treatment_type": treatment.get("treatment_type"),
                    "treatment_or_therapy": treatment.get("treatment_or_therapy"),
                    "therapeutic_agents": treatment.get("therapeutic_agents"),
                    "protocol_identifier": treatment.get("protocol_identifier"),
                    "days_to_treatment_start": treatment.get("days_to_treatment_start"),
                    "days_to_treatment_end": treatment.get("days_to_treatment_end"),
                    "timepoint_category": treatment.get("timepoint_category"),
                    "treatment_outcome": treatment.get("treatment_outcome"),
                    "course_number": treatment.get("course_number"),
                    "payload": dict(treatment),
                }
            )

    follow_ups = [
        {
            **keys,
            "follow_up_id": follow_up.get("follow_up_id"),
            "days_to_follow_up": follow_up.get("days_to_follow_up"),
            "timepoint_category": follow_up.get("timepoint_category"),
            "first_event": follow_up.get("first_event"),
            "days_to_first_event": follow_up.get("days_to_first_event"),
            "year_of_follow_up": follow_up.get("year_of_follow_up"),
            "payload": dict(follow_up),
        }
        for follow_up in as_records(case.get("follow_ups"))
    ]

    return {
        "cases": [case_row],
        "demographics": demographics,
        "diagnoses": diagnoses,
        "follow_ups": follow_ups,
        "treatments": treatments,
    }


def parse_cases(cases: Sequence[Mapping[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    buckets: dict[str, list[dict[str, Any]]] = {
        "cases": [],
        "demographics": [],
        "diagnoses": [],
        "follow_ups": [],
        "treatments": [],
    }
    for case in cases:
        parsed = parse_case_entities(case)
        for name, rows in parsed.items():
            buckets[name].extend(rows)
    return buckets


IDENTIFIER_HEADERS = frozenset({"target usi", "target barcode", "submitter_id"})


def cell_to_jsonable(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, datetime | date):
        return value.isoformat()
    if isinstance(value, str | bool | int | float):
        return value
    return str(value)


def _header_names(row: tuple[Any, ...] | None) -> list[str]:
    if row is None:
        return []
    return [
        f"unnamed_{index + 1}"
        if value is None or str(value).strip() == ""
        else str(value).strip()
        for index, value in enumerate(row)
    ]


def identifier_header(columns: list[str]) -> str | None:
    for name in columns:
        if name.strip().lower() in IDENTIFIER_HEADERS:
            return name
    return None


def _infer_type(values: list[Any]) -> str:
    observed = [value for value in values if is_observed(value)]
    if not observed:
        return "empty"
    kinds: set[str] = set()
    for value in observed:
        if isinstance(value, bool):
            kinds.add("bool")
        elif isinstance(value, int):
            kinds.add("int")
        elif isinstance(value, float):
            kinds.add("float")
        else:
            kinds.add("str")
    if len(kinds) == 1:
        return next(iter(kinds))
    if kinds <= {"int", "float"}:
        return "float"
    return "mixed"


def read_workbook_sheets(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets: list[dict[str, Any]] = []
    try:
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            rows = worksheet.iter_rows(values_only=True)
            columns = _header_names(next(rows, None))
            id_column = identifier_header(columns)
            records: list[dict[str, Any]] = []
            identifiers: list[Any] = []
            for row_number, row in enumerate(rows, start=2):
                cells = {
                    column: cell_to_jsonable(row[index] if index < len(row) else None)
                    for index, column in enumerate(columns)
                }
                identifier = cells.get(id_column) if id_column else None
                identifiers.append(identifier)
                records.append(
                    {
                        "row_number": row_number,
                        "original_identifier": original_identifier(identifier),
                        "normalized_identifier": normalize_identifier(identifier),
                        "join_barcode": join_barcode(identifier),
                        "identifier_shape": identifier_shape(identifier),
                        "cells": cells,
                    }
                )
            sheets.append(
                {
                    "workbook": path.name,
                    "path": str(path),
                    "sheet": sheet_name,
                    "columns": columns,
                    "identifier_field": id_column,
                    "n_rows": len(records),
                    "n_columns": len(columns),
                    "is_patient_level": id_column is not None,
                    "records": records,
                    "identifier_summary": summarize_identifiers(identifiers),
                }
            )
    finally:
        workbook.close()
    return sheets


def profile_sheet(
    sheet: dict[str, Any], *, include_samples: bool = False
) -> list[dict[str, Any]]:
    columns: list[str] = sheet["columns"]
    records: list[dict[str, Any]] = sheet["records"]
    n_rows = len(records)
    id_field = sheet.get("identifier_field")
    profiles: list[dict[str, Any]] = []
    for column in columns:
        values = [record["cells"].get(column) for record in records]
        observed = [value for value in values if is_observed(value)]
        missing = n_rows - len(observed)
        numeric = [
            float(value)
            for value in observed
            if isinstance(value, int | float) and not isinstance(value, bool)
        ]
        profiles.append(
            {
                "workbook": sheet["workbook"],
                "sheet": sheet["sheet"],
                "column": column,
                "is_identifier": column == id_field,
                "n_rows": n_rows,
                "n_missing": missing,
                "pct_missing": round(missing / n_rows * 100.0, 2) if n_rows else 0.0,
                "n_unique_observed": len({repr(value) for value in observed}),
                "inferred_type": _infer_type(values),
                "numeric_min": min(numeric) if numeric else None,
                "numeric_max": max(numeric) if numeric else None,
                "sample_values": (
                    [str(value) for value in observed[:5]]
                    if include_samples and column != id_field
                    else None
                ),
            }
        )
    return profiles


def column_value_counts(sheet: dict[str, Any], column: str) -> Counter[str]:
    counts: Counter[str] = Counter()
    for record in sheet["records"]:
        value = record["cells"].get(column)
        if is_observed(value):
            counts[str(value)] += 1
    return counts

# ---------------------------------------------------------------------------
# Source reconciliation QA
# ---------------------------------------------------------------------------

CYTOGENETIC_COLUMNS: tuple[str, ...] = (
    "t(8;21)",
    "inv(16)",
    "MLL",
    "monosomy 7",
    "monosomy 5",
    "t(6;9)",
    "Primary Cytogenetic Code",
)

DAYS_PER_YEAR = 365.25

BANDS: tuple[tuple[str, float | None, float | None], ...] = (
    ("<1", None, 1.0),
    ("1-4", 1.0, 5.0),
    ("5-9", 5.0, 10.0),
    ("10-14", 10.0, 15.0),
    ("15-17", 15.0, 18.0),
    ("18-21", 18.0, 22.0),
    ("22-29", 22.0, 30.0),
    (">=30", 30.0, None),
)


def days_to_years(days: Any) -> float | None:
    """Convert age-at-diagnosis days to years for QA summaries only."""
    if isinstance(days, bool) or days is None:
        return None
    try:
        number = float(days)
    except (TypeError, ValueError):
        return None
    if number < 0:
        return None
    return number / DAYS_PER_YEAR


def band_for_years(years: float) -> str:
    for label, lower, upper in BANDS:
        if lower is not None and years < lower:
            continue
        if upper is not None and years >= upper:
            continue
        return label
    return "unbanded"


def summarize_age_days(values: Sequence[Any]) -> dict[str, Any]:
    """Summarize age in days for QA without choosing an eligibility cutoff."""
    years_list: list[float] = []
    n_missing = 0
    for value in values:
        years = days_to_years(value)
        if years is None:
            n_missing += 1
        else:
            years_list.append(years)

    bands = {label: 0 for label, *_ in BANDS}
    for years in years_list:
        bands[band_for_years(years)] += 1

    return {
        "n_records": len(values),
        "n_with_age": len(years_list),
        "n_missing_age": n_missing,
        "year_conversion": "days / 365.25",
        "min_years": min(years_list) if years_list else None,
        "max_years": max(years_list) if years_list else None,
        "bands": bands,
        "n_age_lt_18": sum(years < 18 for years in years_list),
        "n_age_le_18": sum(years <= 18 for years in years_list),
        "n_age_le_21": sum(years <= 21 for years in years_list),
        "n_age_ge_18": sum(years >= 18 for years in years_list),
    }


def pairwise_overlap_counts(
    groups: Mapping[str, set[str]],
) -> list[dict[str, str | int]]:
    """Return pairwise intersections, including self-comparisons."""
    names = list(groups)
    rows: list[dict[str, str | int]] = []
    for left in names:
        for right in names:
            rows.append(
                {
                    "file_a": left,
                    "file_b": right,
                    "n_shared": len(groups[left] & groups[right]),
                    "n_a": len(groups[left]),
                    "n_b": len(groups[right]),
                }
            )
    return rows


def overlap_distribution(
    groups: Mapping[str, set[str]],
) -> list[dict[str, int | str]]:
    """Count how many groups each identifier belongs to."""
    membership: Counter[str] = Counter()
    for values in groups.values():
        for item in values:
            membership[item] += 1
    distribution = Counter(membership.values())
    return [
        {"n_files": n_files, "n_identifiers": count}
        for n_files, count in sorted(distribution.items())
    ]


def universe_overlap(left: set[str], right: set[str]) -> dict[str, int | float]:
    """Compare two identifier universes."""
    intersection = left & right
    left_only = left - right
    right_only = right - left
    n_left = len(left)
    n_right = len(right)
    n_intersection = len(intersection)
    return {
        "n_left": n_left,
        "n_right": n_right,
        "n_intersection": n_intersection,
        "n_left_only": len(left_only),
        "n_right_only": len(right_only),
        "pct_left_matched": (
            round(n_intersection / n_left * 100.0, 2) if n_left else 0.0
        ),
        "pct_right_matched": (
            round(n_intersection / n_right * 100.0, 2) if n_right else 0.0
        ),
    }


def _as_number(value: Any) -> float | None:
    if isinstance(value, bool) or not is_observed(value):
        return None
    if isinstance(value, int | float):
        number = float(value)
        return number if number == number else None
    if isinstance(value, str):
        try:
            return float(value.strip().replace(",", ""))
        except ValueError:
            return None
    return None


def _norm_category(value: Any) -> str | None:
    if not is_observed(value):
        return None
    if isinstance(value, str):
        return " ".join(value.strip().lower().split())
    return str(value).strip().lower()


def categorical_agreement(
    pairs: Sequence[tuple[Any, Any]],
) -> dict[str, int | float]:
    """Compare paired categorical values from two sources."""
    both_observed = 0
    agreements = 0
    disagreements = 0
    missing_a_only = 0
    missing_b_only = 0
    missing_both = 0

    for left, right in pairs:
        left_obs = is_observed(left)
        right_obs = is_observed(right)
        if left_obs and right_obs:
            both_observed += 1
            if _norm_category(left) == _norm_category(right):
                agreements += 1
            else:
                disagreements += 1
        elif not left_obs and not right_obs:
            missing_both += 1
        elif not left_obs:
            missing_a_only += 1
        else:
            missing_b_only += 1

    return {
        "n_pairs": len(pairs),
        "n_both_observed": both_observed,
        "n_agreements": agreements,
        "n_disagreements": disagreements,
        "agreement_percent": (
            round(agreements / both_observed * 100.0, 2) if both_observed else 0.0
        ),
        "n_missing_a_only": missing_a_only,
        "n_missing_b_only": missing_b_only,
        "n_missing_both": missing_both,
    }


def _median(values: Sequence[float]) -> float:
    ordered = sorted(values)
    n = len(ordered)
    mid = n // 2
    if n % 2 == 1:
        return float(ordered[mid])
    return (ordered[mid - 1] + ordered[mid]) / 2.0


def numeric_discordance(pairs: Sequence[tuple[Any, Any]]) -> dict[str, Any]:
    """Compare paired numeric values, including difference summaries."""
    both_observed = 0
    exact = 0
    within_one = 0
    missing_a_only = 0
    missing_b_only = 0
    missing_both = 0
    diffs: list[float] = []

    for left, right in pairs:
        left_n = _as_number(left)
        right_n = _as_number(right)
        if left_n is not None and right_n is not None:
            both_observed += 1
            diff = left_n - right_n
            diffs.append(diff)
            if diff == 0:
                exact += 1
            if abs(diff) <= 1:
                within_one += 1
        elif left_n is None and right_n is None:
            missing_both += 1
        elif left_n is None:
            missing_a_only += 1
        else:
            missing_b_only += 1

    abs_diffs = [abs(item) for item in diffs]
    return {
        "n_pairs": len(pairs),
        "n_both_observed": both_observed,
        "n_exact_agreements": exact,
        "n_within_one": within_one,
        "agreement_percent": (
            round(exact / both_observed * 100.0, 2) if both_observed else 0.0
        ),
        "n_disagreements": both_observed - exact,
        "n_missing_a_only": missing_a_only,
        "n_missing_b_only": missing_b_only,
        "n_missing_both": missing_both,
        "diff_min": min(diffs) if diffs else None,
        "diff_max": max(diffs) if diffs else None,
        "abs_diff_median": _median(abs_diffs) if abs_diffs else None,
        "n_nonzero_diffs": sum(diff != 0 for diff in diffs),
    }


DEFAULT_OUTPUT = PROJECT_ROOT / "artifacts" / "ingestion_audit"
DETAIL_DIR = PROJECT_ROOT / "data" / "interim" / "ingestion_audit"

NUMERIC_SUPPLEMENT_FIELDS = {
    "wbc": "wbc_raw",
    "marrow_blasts": "marrow_blasts_raw",
    "peripheral_blasts": "peripheral_blasts_raw",
}
CATEGORICAL_SUPPLEMENT_FIELDS = {
    "risk_group": "risk_group_raw",
    "flt3_itd": "flt3_itd_raw",
    "npm": "npm_raw",
    "cebpa": "cebpa_raw",
    "fab": "fab_raw",
    "cns_disease": "cns_disease_raw",
}

_OBSOLETE_ARTIFACTS = (
    "age_distribution.json",
    "clinical_concept_source_map.csv",
    "entity_counts.json",
    "missing_value_policy.json",
    "missing_value_token_inventory.csv",
    "supplement_identifier_quality.csv",
    "supplement_overlap_distribution.csv",
    "supplement_overlap_matrix.csv",
    "supplement_sheet_summary.csv",
    "unmatched_identifier_summary.csv",
)


def _write(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.suffix == ".csv":
        rows = payload
        if not rows:
            path.write_text("", encoding="utf-8")
            return
        with path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)
        return
    path.write_text(
        json.dumps(payload, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def _cells(row: dict[str, Any]) -> dict[str, Any]:
    value = row.get("cells")
    if isinstance(value, dict):
        return value
    return json.loads(value) if isinstance(value, str) else {}


def _fetch_reconciliation_rows(
    engine: Engine,
) -> dict[str, list[dict[str, Any]]]:
    queries = {
        "cases": """
            SELECT case_id, submitter_id, join_barcode
            FROM staging.gdc_cases
        """,
        "demographics": "SELECT * FROM staging.gdc_demographics",
        "diagnoses": "SELECT * FROM staging.gdc_diagnoses",
        "follow_ups": "SELECT * FROM staging.gdc_follow_ups",
        "supplements": """
            SELECT * FROM staging.supplement_clinical_rows
            WHERE sheet_name IN ('Clinical Data', 'Sheet1')
        """,
    }
    with engine.connect() as connection:
        return {
            name: [
                dict(row)
                for row in connection.execute(text(sql)).mappings()
            ]
            for name, sql in queries.items()
        }


def _supplements_by_workbook(
    rows: list[dict[str, Any]],
) -> dict[str, dict[str, dict[str, Any]]]:
    out: dict[str, dict[str, dict[str, Any]]] = defaultdict(dict)
    for row in rows:
        if barcode := row.get("join_barcode"):
            out[row["workbook_name"]][barcode] = row
    return out


def _identifier_overlap(
    cases: list[dict[str, Any]],
    supplements: dict[str, dict[str, dict[str, Any]]],
) -> dict[str, Any]:
    gdc_ids = {row["join_barcode"] for row in cases if row["join_barcode"]}
    supplement_ids = (
        set().union(*(set(rows) for rows in supplements.values()))
        if supplements
        else set()
    )
    stats = universe_overlap(gdc_ids, supplement_ids)
    return {
        "gdc_unique_join_barcodes": stats["n_left"],
        "supplement_unique_join_barcodes": stats["n_right"],
        "intersection": stats["n_intersection"],
        "gdc_only": stats["n_left_only"],
        "supplement_only": stats["n_right_only"],
        "pct_gdc_matched": stats["pct_left_matched"],
        "pct_supplement_matched": stats["pct_right_matched"],
        "normalization_rule": (
            "strip whitespace, uppercase; join_barcode is leading TARGET-NN-TOKEN"
        ),
    }


def _gdc_index(data: dict[str, list[dict[str, Any]]]) -> dict[str, dict[str, Any]]:
    demographics = {
        row["join_barcode"]: row
        for row in data["demographics"]
        if row["join_barcode"]
    }
    diagnoses = {
        row["join_barcode"]: row
        for row in data["diagnoses"]
        if row["join_barcode"]
    }

    last_contact: dict[str, list[float]] = defaultdict(list)
    for row in data["follow_ups"]:
        barcode = row["join_barcode"]
        value = row["days_to_follow_up"]
        if (
            barcode
            and value is not None
            and row["timepoint_category"] == "Last Contact"
        ):
            last_contact[barcode].append(float(value))

    out: dict[str, dict[str, Any]] = {}
    for barcode, demographic in demographics.items():
        diagnosis = diagnoses.get(barcode, {})
        status = demographic["vital_status_analysis_class"]
        death = demographic["days_to_death"]
        follow = diagnosis.get("days_to_last_follow_up")
        os_days = None
        if status == "dead" and death is not None:
            os_days = float(death)
        elif status == "alive" and follow is not None:
            os_days = float(follow)

        out[barcode] = {
            "vital_status": demographic["vital_status_raw"],
            "os_days": os_days,
            "days_to_death": death,
            "days_to_last_follow_up": follow,
            "last_contact_days": (
                max(last_contact[barcode]) if last_contact[barcode] else None
            ),
            "age_at_diagnosis_days": diagnosis.get("age_at_diagnosis_days"),
            "sex_at_birth": demographic["sex_at_birth_raw"],
        }
    return out


def _compact_stats(
    workbook: str,
    concept: str,
    stats: dict[str, Any],
    *,
    numeric: bool,
) -> dict[str, Any]:
    return {
        "workbook": workbook,
        "concept": concept,
        "n_both_observed": stats["n_both_observed"],
        "n_agreements": (
            stats["n_exact_agreements"] if numeric else stats["n_agreements"]
        ),
        "n_disagreements": stats["n_disagreements"],
        "agreement_percent": stats["agreement_percent"],
        "n_missing_a_only": stats["n_missing_a_only"],
        "n_missing_b_only": stats["n_missing_b_only"],
        "n_missing_both": stats["n_missing_both"],
    }


def _gdc_vs_supplement(
    data: dict[str, list[dict[str, Any]]],
    supplements: dict[str, dict[str, dict[str, Any]]],
) -> tuple[
    list[dict[str, Any]],
    list[dict[str, Any]],
    list[dict[str, Any]],
]:
    gdc = _gdc_index(data)
    os_rows: list[dict[str, Any]] = []
    summary_rows: list[dict[str, Any]] = []
    detail_rows: list[dict[str, Any]] = []

    for workbook, records in supplements.items():
        shared = set(records) & set(gdc)
        status_pairs = []
        time_pairs = []
        contact_pairs = []
        age_pairs = []
        sex_pairs = []

        for barcode in shared:
            left = gdc[barcode]
            right = records[barcode]
            status_pair = (left["vital_status"], right.get("vital_status_raw"))
            time_pair = (left["os_days"], right.get("os_time_days"))
            status_pairs.append(status_pair)
            time_pairs.append(time_pair)
            contact_pairs.append(
                (left["last_contact_days"], right.get("os_time_days"))
            )
            age_pairs.append(
                (
                    left["age_at_diagnosis_days"],
                    right.get("age_at_diagnosis_days"),
                )
            )
            sex_pairs.append((left["sex_at_birth"], right.get("sex_raw")))

            status_one = categorical_agreement([status_pair])
            time_one = numeric_discordance([time_pair])
            if status_one["n_disagreements"] or time_one["n_disagreements"]:
                detail_rows.append(
                    {
                        "join_barcode": barcode,
                        "workbook": workbook,
                        "gdc_vital_status": left["vital_status"],
                        "supplement_vital_status": right.get("vital_status_raw"),
                        "gdc_candidate_os_days": left["os_days"],
                        "gdc_days_to_death": left["days_to_death"],
                        "gdc_days_to_last_follow_up": left[
                            "days_to_last_follow_up"
                        ],
                        "gdc_last_contact_days": left["last_contact_days"],
                        "supplement_os_time_days": right.get("os_time_days"),
                    }
                )

        status = categorical_agreement(status_pairs)
        times = numeric_discordance(time_pairs)
        contact = numeric_discordance(contact_pairs)
        ages = numeric_discordance(age_pairs)
        sex = categorical_agreement(sex_pairs)

        patterns = Counter(
            f"{str(left).strip().lower()}|{str(right).strip().lower()}"
            for left, right in status_pairs
        )
        supplement_times = [
            records[barcode].get("os_time_days") for barcode in shared
        ]
        gdc_times = [gdc[barcode]["os_days"] for barcode in shared]

        os_rows.append(
            {
                "comparison": f"gdc_vs_{workbook}",
                "n_shared_patients": len(shared),
                "vital_status_both_observed": status["n_both_observed"],
                "vital_status_agreements": status["n_agreements"],
                "vital_status_disagreements": status["n_disagreements"],
                "vital_status_agreement_percent": status["agreement_percent"],
                "vital_status_pair_patterns": json.dumps(
                    dict(patterns), sort_keys=True
                ),
                "os_time_both_observed": times["n_both_observed"],
                "os_time_exact_agreements": times["n_exact_agreements"],
                "os_time_within_one_day": times["n_within_one"],
                "os_time_agreement_percent": times["agreement_percent"],
                "os_time_diff_min": times["diff_min"],
                "os_time_diff_max": times["diff_max"],
                "os_time_abs_diff_median": times["abs_diff_median"],
                "last_contact_vs_os_exact": contact["n_exact_agreements"],
                "last_contact_vs_os_both_observed": contact["n_both_observed"],
                "last_contact_vs_os_agreement_percent": contact[
                    "agreement_percent"
                ],
                "age_exact_agreements": ages["n_exact_agreements"],
                "age_both_observed": ages["n_both_observed"],
                "age_agreement_percent": ages["agreement_percent"],
                "sex_agreement_percent": sex["agreement_percent"],
                "sex_disagreements": sex["n_disagreements"],
                "supplement_zero_os_times_in_overlap": sum(
                    value == 0 for value in supplement_times
                ),
                "supplement_negative_os_times_in_overlap": sum(
                    value is not None and value < 0 for value in supplement_times
                ),
                "gdc_zero_candidate_os_in_overlap": sum(
                    value == 0 for value in gdc_times
                ),
                "gdc_negative_candidate_os_in_overlap": sum(
                    value is not None and value < 0 for value in gdc_times
                ),
            }
        )
        summary_rows.extend(
            (
                _compact_stats(
                    workbook, "vital_status", status, numeric=False
                ),
                _compact_stats(
                    workbook,
                    "os_time_gdc_candidate_vs_supplement",
                    times,
                    numeric=True,
                ),
                _compact_stats(
                    workbook, "age_at_diagnosis", ages, numeric=True
                ),
                _compact_stats(
                    workbook,
                    "sex_gdc_sex_at_birth_vs_supplement_gender",
                    sex,
                    numeric=False,
                ),
            )
        )

    return os_rows, summary_rows, detail_rows


def _supplement_discordance(
    supplements: dict[str, dict[str, dict[str, Any]]],
) -> list[dict[str, Any]]:
    comparisons = [
        *(
            (name, "staging", column, True)
            for name, column in NUMERIC_SUPPLEMENT_FIELDS.items()
        ),
        *(
            (name, "staging", column, False)
            for name, column in CATEGORICAL_SUPPLEMENT_FIELDS.items()
        ),
        *(
            (f"cytogenetics:{column}", "cells", column, False)
            for column in CYTOGENETIC_COLUMNS
        ),
    ]
    rows: list[dict[str, Any]] = []
    workbooks = sorted(supplements)

    for index, left_name in enumerate(workbooks):
        for right_name in workbooks[index + 1 :]:
            shared = set(supplements[left_name]) & set(supplements[right_name])
            for concept, origin, column, numeric in comparisons:
                pairs = []
                for barcode in shared:
                    left = supplements[left_name][barcode]
                    right = supplements[right_name][barcode]
                    if origin == "cells":
                        pairs.append(
                            (
                                _cells(left).get(column),
                                _cells(right).get(column),
                            )
                        )
                    else:
                        pairs.append((left.get(column), right.get(column)))

                stats = (
                    numeric_discordance(pairs)
                    if numeric
                    else categorical_agreement(pairs)
                )
                rows.append(
                    {
                        "file_a": left_name,
                        "file_b": right_name,
                        "concept": concept,
                        "value_kind": "numeric" if numeric else "categorical",
                        "n_shared_patients": len(shared),
                        "n_both_observed": stats["n_both_observed"],
                        "n_agreements": (
                            stats["n_exact_agreements"]
                            if numeric
                            else stats["n_agreements"]
                        ),
                        "n_disagreements": stats["n_disagreements"],
                        "agreement_percent": stats["agreement_percent"],
                        "n_missing_a_only": stats["n_missing_a_only"],
                        "n_missing_b_only": stats["n_missing_b_only"],
                        "n_missing_both": stats["n_missing_both"],
                        "diff_min": stats["diff_min"] if numeric else None,
                        "diff_max": stats["diff_max"] if numeric else None,
                        "abs_diff_median": (
                            stats["abs_diff_median"] if numeric else None
                        ),
                    }
                )
    return rows


def _source_decisions() -> dict[str, Any]:
    return {
        "overall_survival": {
            "source": "GDC",
            "event": "demographic.vital_status",
            "time": (
                "days_to_death for deaths; diagnoses.days_to_last_follow_up "
                "for survivors"
            ),
            "reason": (
                "Overlapping supplement OS fields are retained for QA but "
                "are not treated as the primary endpoint source."
            ),
        },
        "age_at_diagnosis": {
            "source": "GDC diagnoses.age_at_diagnosis",
            "reason": "Used for cohort eligibility and baseline age.",
        },
        "sex": {
            "source": "GDC demographic.sex_at_birth",
            "reason": "The supplement Gender field is not substituted.",
        },
        "aml_specific_baselines": {
            "source": "TARGET clinical supplements",
            "reason": (
                "WBC, protocol risk, molecular markers, morphology, and "
                "cytogenetic fields are primarily available in supplements."
            ),
            "precedence_note": (
                "Per-variable workbook precedence is locked in "
                "cohort/baseline.py and is not outcome-driven."
            ),
        },
    }


def _compatibility_summaries(
    data: dict[str, list[dict[str, Any]]],
    supplements: dict[str, dict[str, dict[str, Any]]],
) -> dict[str, Any]:
    supplement_sets = {
        workbook: set(rows) for workbook, rows in supplements.items()
    }
    return {
        "entity_counts": {},
        "age": summarize_age_days(
            [row["age_at_diagnosis_days"] for row in data["diagnoses"]]
        ),
        "gdc_id_summary": summarize_identifiers(
            [row["submitter_id"] for row in data["cases"]]
        ),
        "supp_id_summary": summarize_identifiers(
            [row["original_identifier"] for row in data["supplements"]]
        ),
        "overlap_distribution": overlap_distribution(supplement_sets),
        "pairwise_overlap": pairwise_overlap_counts(supplement_sets),
    }


def _remove_obsolete(output_dir: Path) -> None:
    for filename in _OBSOLETE_ARTIFACTS:
        (output_dir / filename).unlink(missing_ok=True)


def run_reconciliation(
    engine: Engine,
    output_dir: Path = DEFAULT_OUTPUT,
    detail_dir: Path = DETAIL_DIR,
) -> dict[str, Any]:
    """Validate source linkage and concordance before cohort construction."""
    data = _fetch_reconciliation_rows(engine)
    supplements = _supplements_by_workbook(data["supplements"])

    identifier_overlap = _identifier_overlap(data["cases"], supplements)
    os_rows, source_discordance, os_detail = _gdc_vs_supplement(
        data, supplements
    )
    supplement_discordance = _supplement_discordance(supplements)
    decisions = _source_decisions()

    _remove_obsolete(output_dir)
    _write(output_dir / "source_identifier_overlap.json", identifier_overlap)
    _write(output_dir / "os_source_reconciliation.csv", os_rows)
    _write(
        output_dir / "gdc_vs_supplement_discordance.csv",
        source_discordance,
    )
    _write(
        output_dir / "supplement_discordance_summary.csv",
        supplement_discordance,
    )
    _write(output_dir / "source_decisions.json", decisions)
    if os_detail:
        _write(detail_dir / "os_discordance_examples.csv", os_detail)

    compatibility = _compatibility_summaries(data, supplements)
    return {
        "identifier_overlap": identifier_overlap,
        "entity_counts": compatibility["entity_counts"],
        "age": compatibility["age"],
        "os_rows": os_rows,
        "discordance_rows": supplement_discordance,
        "gdc_id_summary": compatibility["gdc_id_summary"],
        "supp_id_summary": compatibility["supp_id_summary"],
        "sheet_summary": [],
        "overlap_distribution": compatibility["overlap_distribution"],
        "pairwise_overlap": compatibility["pairwise_overlap"],
        "source_decisions": decisions,
    }

